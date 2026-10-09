"""Output side: DOCX assembly on top of a company template (python-docx)."""
import copy
import csv
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .sources import Source

# Table look copied from technicka_sprava_sablona.docx
HEADER_FILL = "EEF3F8"
BORDER = "CCCCCC"
GREY = RGBColor(0x80, 0x80, 0x80)
BODY_BOOKMARK = "RB_Body"


def _slug(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"[^0-9A-Za-z]+", "_", text).strip("_")
    return text.lower()[:60] or "table"


class DocBuilder:
    """Fills the template's title page, drops its sample body and appends content.

    Every table is also saved as CSV and every figure stays as PNG, so the
    output folder doubles as raw material for a hand-written report.
    """

    def __init__(self, out_dir: Path, template: Optional[Path], title: str, subtitle: str,
                 campaign: str, author: str = "", status: str = "Draft"):
        self.out_dir = Path(out_dir)
        self.tables_dir = self.out_dir / "tables"
        self.figures_dir = self.out_dir / "figures"
        for d in (self.out_dir, self.tables_dir, self.figures_dir):
            d.mkdir(parents=True, exist_ok=True)
        self._h = [0, 0, 0]
        self._n_tab = 0
        self._n_fig = 0
        self._captions = []      # (paragraph, kind, text) - numbered in document order at save()
        self._toc = []          # (level, text, bookmark) of every heading we add
        self._toc_par = {}      # bookmark -> heading paragraph element
        self._toc_anchor = None  # element the table of contents is written after
        self._date = datetime.now().strftime("%d.%m.%Y")
        self.doc = Document(str(template)) if template else Document()
        if template:
            self._prepare_template(title, subtitle, campaign, author, status)
        else:
            self.doc.add_heading(title, 0)
            self.doc.add_paragraph(f"{subtitle} | Kampaň: {campaign} | {self._date} | {author}")

    # ------------------------------------------------------------------ template
    def _prepare_template(self, title, subtitle, campaign, author, status) -> None:
        doc = self.doc
        fill = {
            "[NAZOV KAMPANE]": campaign, "[NAZOV]": campaign, "[DD.MM.RRRR]": self._date,
            "[Meno Priezvisko]": author or "[Meno Priezvisko]", "[Meno]": author or "[Meno]",
            "Draft / Final": status,
        }
        # Title page: the 2nd non-empty paragraph is the subtitle line.
        seen = 0
        for p in doc.paragraphs:
            if p.style.name.startswith("Heading"):
                break
            if p.text.strip():
                seen += 1
                if seen == 2:
                    self._set_text(p, subtitle)
                    break
        paragraphs = list(doc.paragraphs)
        for t in doc.tables:
            paragraphs += [p for row in t.rows for c in row.cells for p in c.paragraphs]
        for s in doc.sections:
            for p in s.header.paragraphs:
                if "Kampaň:" in p.text:
                    self._set_text(p, f"Technická správa – {title}  |  Kampaň: {campaign}")
        for p in paragraphs:
            for old, new in fill.items():
                if old in p.text:
                    self._set_text(p, p.text.replace(old, new))
        self._fix_page_number()
        self._cut_sample_body()

    def _fix_page_number(self) -> None:
        """The template footer holds the literal text 'PAGE'; turn it into a real page field."""
        for section in self.doc.sections:
            for p in section.footer.paragraphs:
                for r in p.runs:
                    if not r.text.endswith("PAGE"):
                        continue
                    r.text = r.text[:-4]
                    last = r._r
                    for kind, instr in (("begin", None), (None, " PAGE "), ("end", None)):
                        new = copy.deepcopy(r._r)
                        for child in list(new):
                            if child.tag != qn("w:rPr"):
                                new.remove(child)
                        if kind:
                            el = OxmlElement("w:fldChar")
                            el.set(qn("w:fldCharType"), kind)
                        else:
                            el = OxmlElement("w:instrText")
                            el.set(qn("xml:space"), "preserve")
                            el.text = instr
                        new.append(el)
                        last.addnext(new)
                        last = new
                    return

    @staticmethod
    def _set_text(paragraph, text: str) -> None:
        runs = paragraph.runs
        if not runs:
            paragraph.add_run(text)
            return
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""

    def _cut_sample_body(self) -> None:
        """Remove everything after the table of contents (the template's sample chapters)."""
        body = self.doc.element.body
        children = list(body.iterchildren())
        toc = next((i for i, el in enumerate(children) if el.tag == qn("w:sdt")), None)
        if toc is None:
            return
        for el in children[toc + 1:]:
            if el.tag != qn("w:sectPr"):
                body.remove(el)
        # The template's own contents list points at the removed chapters; ours is
        # written in save(), after the heading that precedes it.
        self._toc_anchor = children[toc - 1] if toc else None
        body.remove(children[toc])
        self.page_break()

    # ------------------------------------------------------------------ text
    def heading(self, text: str, level: int = 1, numbered: bool = True) -> None:
        if numbered:
            self._h[level - 1] += 1
            for i in range(level, 3):
                self._h[i] = 0
            number = ".".join(str(n) for n in self._h[:level])
            text = f"{number}  {text}"
        p = self.doc.add_heading(text, level)
        name = f"_TocRB{len(self._toc) + 1:04d}"
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        for el in (start, end):
            el.set(qn("w:id"), str(9000 + len(self._toc)))
        start.set(qn("w:name"), name)
        p._p.insert(1 if p._p.pPr is not None else 0, start)
        p._p.append(end)
        self._toc.append((level, text, name))
        self._toc_par[name] = p._p

    def para(self, text: str, italic: bool = False, size: Optional[float] = None):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        return p

    def comment(self, hint: str = "") -> None:
        """Visible slot for the author's own commentary."""
        p = self.doc.add_paragraph()
        run = p.add_run(f"[Komentár: {hint}]" if hint else "[Komentár]")
        run.italic = True
        run.font.color.rgb = GREY

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def landscape(self, on: bool = True) -> None:
        """Start a new section in landscape (wide tables) or back in portrait."""
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        want = WD_ORIENT.LANDSCAPE if on else WD_ORIENT.PORTRAIT
        if section.orientation != want:
            section.orientation = want
            section.page_width, section.page_height = section.page_height, section.page_width

    def _caption(self, text: str) -> None:
        p = self.doc.add_paragraph()
        self._captions.append((p, "Tabuľka", text))
        run = p.add_run(text)
        run.italic = True
        run.font.size = Pt(9)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)

    def _text_width_cm(self) -> float:
        s = self.doc.sections[-1]
        return (s.page_width - s.left_margin - s.right_margin) / 360000.0

    # ------------------------------------------------------------------ table
    def table(self, caption: str, columns: Sequence[str], rows: Sequence[Sequence],
              align: Optional[str] = None, name: Optional[str] = None) -> None:
        """Add a captioned table and save it as CSV.

        align: one letter per column, 'l' or 'r'. Default: first column left, rest right.
        """
        self._n_tab += 1
        rows = [[str(v) for v in row] for row in rows]
        self._write_csv(name or f"{self._n_tab:02d}_{_slug(caption)}", columns, rows)
        self._caption(caption)
        n = len(columns)
        align = align or ("l" + "r" * (n - 1))
        size = 9 if n <= 6 else 8 if n <= 9 else 7 if n <= 13 else 6.5

        table = self.doc.add_table(rows=len(rows) + 1, cols=n)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False  # fixed layout: Word and LibreOffice keep our widths
        self._borders(table)
        # Column widths proportional to the longest cell, within the text column.
        # A header may wrap between words, a data cell should not: weigh each column by
        # its longest header word and its longest cell.
        longest = []
        for i in range(n):
            # header is bold, so its words need a bit more room than data of equal length
            words = max(len(w) for w in str(columns[i]).split() or [""]) * 1.25
            longest.append(min(max([words] + [len(r[i]) for r in rows]), 45) + 4)
        total = self._text_width_cm()
        widths = [Cm(total * v / sum(longest)) for v in longest]
        for col, width in zip(table.columns, widths):
            col.width = width

        cells = table._cells  # one lookup; cell-by-cell access is slow on big tables
        for idx, cell in enumerate(cells):
            r, c = divmod(idx, n)
            text = str(columns[c]) if r == 0 else rows[r - 1][c]
            cell.width = widths[c]
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(size)
            p.paragraph_format.space_after = Pt(0)
            if r == 0:
                run.bold = True
                self._shade(cell, HEADER_FILL)
            if align[c] == "r":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._repeat_header(table.rows[0])
        self.doc.add_paragraph()

    def _write_csv(self, name: str, columns, rows) -> None:
        # UTF-8 BOM + ';' so Excel opens it directly.
        with open(self.tables_dir / f"{name}.csv", "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(columns)
            writer.writerows(rows)

    @staticmethod
    def _borders(table) -> None:
        pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), BORDER)
            borders.append(el)
        pr.append(borders)
        margins = OxmlElement("w:tblCellMar")
        for edge in ("left", "right"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:w"), "60")
            el.set(qn("w:type"), "dxa")
            margins.append(el)
        pr.append(margins)

    @staticmethod
    def _shade(cell, fill: str) -> None:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shd)

    @staticmethod
    def _repeat_header(row) -> None:
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(el)

    def key_value(self, caption: str, pairs: Sequence[Sequence], name: Optional[str] = None) -> None:
        self.table(caption, ["Parameter", "Hodnota"], pairs, align="ll", name=name)

    # ------------------------------------------------------------------ figure
    def figure(self, caption: str, png: Path, width_cm: Optional[float] = None) -> None:
        self._n_fig += 1
        width = min(width_cm or 99, self._text_width_cm())
        self.doc.add_picture(str(png), width=Cm(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = self.doc.add_paragraph()
        self._captions.append((p, "Obrázok", caption))
        run = p.add_run(caption)
        run.italic = True
        run.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def fig_path(self, name: str) -> Path:
        return self.figures_dir / f"{name}.png"

    # ------------------------------------------------------------------ closing
    def sources(self, items: List[Source]) -> None:
        """Which files the report was built from and when each of them was written."""
        self.table("Zdrojové súbory správy", ["Súbor", "Vytvorený", "Poznámka"],
                   [[s.path.name, s.generated or "-", s.note] for s in items],
                   align="lll", name="00_sources")

    def _write_toc(self) -> None:
        """Table of contents with live page references (PAGEREF to heading bookmarks).

        Page numbers are computed by the word processor itself, so they are right in
        both Word and LibreOffice without a manual field update.
        """
        if self._toc_anchor is None:
            return
        # Word shows the cached "?" until fields are refreshed; ask it to do so on open.
        settings = self.doc.settings.element
        if settings.find(qn("w:updateFields")) is None:
            flag = OxmlElement("w:updateFields")
            flag.set(qn("w:val"), "true")
            settings.append(flag)
        right = int(self._text_width_cm() * 567)  # cm -> twips
        anchor = self._toc_anchor
        entries = [e for e in self._toc if e[0] <= 2]
        if not entries:
            return
        self._bookmark_body()

        def field_run(fld=None, instr=None):
            r = OxmlElement("w:r")
            if fld:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), fld)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = instr
            r.append(el)
            return r

        for index, (level, text, name) in enumerate(entries):
            p = OxmlElement("w:p")
            ppr = OxmlElement("w:pPr")
            tabs, tab = OxmlElement("w:tabs"), OxmlElement("w:tab")
            tab.set(qn("w:val"), "right")
            tab.set(qn("w:leader"), "dot")
            tab.set(qn("w:pos"), str(right))
            tabs.append(tab)
            ind, spacing = OxmlElement("w:ind"), OxmlElement("w:spacing")
            ind.set(qn("w:left"), str(280 * (level - 1)))
            spacing.set(qn("w:after"), "40")
            for el in (tabs, spacing, ind):
                ppr.append(el)
            p.append(ppr)
            if index == 0:
                # A real TOC field around the entries: Word rebuilds it itself (with correct
                # page numbers) when it refreshes fields; the entries below are what
                # LibreOffice shows, with page references it computes live.
                p.append(field_run(fld="begin"))
                p.append(field_run(instr=f' TOC \\o "1-2" \\h \\z \\b {BODY_BOOKMARK} '))
                p.append(field_run(fld="separate"))
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), name)
            link.set(qn("w:history"), "1")

            def run(text=None, fld=None, instr=None, tab_char=False):
                r = OxmlElement("w:r")
                if level == 1:
                    rpr = OxmlElement("w:rPr")
                    rpr.append(OxmlElement("w:b"))
                    r.append(rpr)
                if fld:
                    el = OxmlElement("w:fldChar")
                    el.set(qn("w:fldCharType"), fld)
                elif instr:
                    el = OxmlElement("w:instrText")
                    el.set(qn("xml:space"), "preserve")
                    el.text = instr
                elif tab_char:
                    el = OxmlElement("w:tab")
                else:
                    el = OxmlElement("w:t")
                    el.set(qn("xml:space"), "preserve")
                    el.text = text
                r.append(el)
                link.append(r)

            run(text)
            run(tab_char=True)
            run(fld="begin")
            run(instr=f" PAGEREF {name} \\h ")
            run(fld="separate")
            run("?")
            run(fld="end")
            p.append(link)
            if index == len(entries) - 1:
                p.append(field_run(fld="end"))
            anchor.addnext(p)
            anchor = p

    def _bookmark_body(self) -> None:
        """Bookmark everything from the first report heading to the end, so that the
        contents list covers the report only (not the title page or 'Obsah' itself)."""
        body = self.doc.element.body
        order = {id(el): i for i, el in enumerate(body.iter(qn("w:p")))}
        first = min(self._toc_par.values(), key=lambda el: order.get(id(el), 10 ** 9))
        last = [el for el in body.iterchildren(qn("w:p"))][-1]
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        for el in (start, end):
            el.set(qn("w:id"), "8999")
        start.set(qn("w:name"), BODY_BOOKMARK)
        first.insert(1 if first.pPr is not None else 0, start)
        last.append(end)

    def _number_captions(self) -> None:
        """Number tables and figures in reading order (chapters may have been moved)."""
        order = {id(p): i for i, p in enumerate(self.doc.element.body.iter(qn("w:p")))}
        counters = {}
        for p, kind, text in sorted(self._captions, key=lambda c: order.get(id(c[0]._p), 0)):
            counters[kind] = counters.get(kind, 0) + 1
            p.runs[0].text = f"{kind} {counters[kind]} – {text}"

    def save(self, filename: str) -> Path:
        self._number_captions()
        self._write_toc()
        path = self.out_dir / filename
        self.doc.save(str(path))
        return path
